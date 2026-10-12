import pandas as pd
import numpy as np
from pathlib import Path

project_root = Path('../').resolve()

# ============================================================
# 1. SETTINGS
# ============================================================

INPUT_FILE = project_root/ 'final_roselle' / 'data' / 'final_dataset' / 'final_roselle_dataset.csv'
OUTPUT_FILE = Path("roselle_analog_descriptive_statistics.xlsx")

# Columns that identify compounds rather than descriptors
ID_COLS = [
    "CID",
    "MolecularFormula",
    "smiles",
    "ConnectivitySMILES",
    "InChIKey",
    "IUPACName"
]

# ============================================================
# 2. LOAD DATA
# ============================================================

df = pd.read_csv(INPUT_FILE)

print("=" * 70)
print("ROSELLE ANALOG LIBRARY")
print("=" * 70)

print(f"Number of compounds: {len(df):,}")
print(f"Number of columns:   {len(df.columns):,}")

# Remove accidental whitespace from column names
df.columns = df.columns.str.strip()

# ============================================================
# 3. DEFINE RESEARCH VARIABLES
# ============================================================

# ------------------------------------------------------------
# Physicochemical descriptors
# ------------------------------------------------------------

PHYSICOCHEMICAL_COLS = [
    "MolecularWeight",
    "mol_weight",
    "logp",
    "tpsa",
    "num_h_donors",
    "num_h_acceptors",
    "num_rotatable_bonds",
    "num_aromatic_rings",
    "num_rings",
    "num_heavy_atoms",
    "fraction_csp3",
    "num_valence_electrons",
    "molar_refractivity",
    "num_heteroatoms",
    "num_saturated_rings",
    "balaban_j",
    "qed"
]

# ------------------------------------------------------------
# Predicted bioactivity
# ------------------------------------------------------------

BIOACTIVITY_COLS = [
    "pIC50 against EGFR",
    "pIC50 against HER2",
    "pIC50 against VEGFR2"
]

# ------------------------------------------------------------
# Predicted ADMET
# ------------------------------------------------------------

ADMET_COLS = [
    "mutagenicity",
    "herg",
    "cyp3a4",
    "solubility"
]

# ============================================================
# 4. CHECK WHICH COLUMNS EXIST
# ============================================================

all_requested = (
    PHYSICOCHEMICAL_COLS +
    BIOACTIVITY_COLS +
    ADMET_COLS
)

missing_columns = [
    col for col in all_requested
    if col not in df.columns
]

if missing_columns:
    print("\nWARNING: The following requested columns were not found:")
    for col in missing_columns:
        print(f"  - {col}")

# Keep only columns actually present
PHYSICOCHEMICAL_COLS = [
    col for col in PHYSICOCHEMICAL_COLS
    if col in df.columns
]

BIOACTIVITY_COLS = [
    col for col in BIOACTIVITY_COLS
    if col in df.columns
]

ADMET_COLS = [
    col for col in ADMET_COLS
    if col in df.columns
]

# ============================================================
# 5. CONVERT NUMERIC VARIABLES
# ============================================================

CONTINUOUS_COLS = (
    PHYSICOCHEMICAL_COLS +
    BIOACTIVITY_COLS
)

for col in CONTINUOUS_COLS:
    df[col] = pd.to_numeric(df[col], errors="coerce")

# ============================================================
# 6. CONTINUOUS DESCRIPTIVE STATISTICS
# ============================================================

def continuous_statistics(data, columns):
    """
    Generate comprehensive descriptive statistics
    for continuous variables.
    """

    results = []

    for col in columns:

        x = data[col].dropna()

        n = len(x)
        missing = data[col].isna().sum()

        if n == 0:
            results.append({
                "Variable": col,
                "N": 0,
                "Missing": missing,
                "Mean": np.nan,
                "SD": np.nan,
                "Median": np.nan,
                "Q1 (25%)": np.nan,
                "Q3 (75%)": np.nan,
                "IQR": np.nan,
                "Minimum": np.nan,
                "Maximum": np.nan,
                "Skewness": np.nan,
                "Kurtosis": np.nan
            })
            continue

        q1 = x.quantile(0.25)
        q3 = x.quantile(0.75)

        results.append({
            "Variable": col,
            "N": n,
            "Missing": missing,
            "Mean": x.mean(),
            "SD": x.std(),
            "Median": x.median(),
            "Q1 (25%)": q1,
            "Q3 (75%)": q3,
            "IQR": q3 - q1,
            "Minimum": x.min(),
            "Maximum": x.max(),
            "Skewness": x.skew(),
            "Kurtosis": x.kurtosis()
        })

    return pd.DataFrame(results)


physchem_stats = continuous_statistics(
    df,
    PHYSICOCHEMICAL_COLS
)

bioactivity_stats = continuous_statistics(
    df,
    BIOACTIVITY_COLS
)

# ============================================================
# 7. ADMET DESCRIPTIVE STATISTICS
# ============================================================

def admet_statistics(data, columns):
    """
    Generate descriptive statistics suitable for
    categorical/binary ADMET predictions.

    If an ADMET variable is numeric/continuous,
    additional numerical statistics are also reported.
    """

    results = []

    for col in columns:

        series = data[col]
        non_missing = series.dropna()

        n = len(non_missing)
        missing = series.isna().sum()
        unique = non_missing.nunique()

        # Determine whether the variable is numeric
        numeric = pd.to_numeric(
            non_missing,
            errors="coerce"
        )

        numeric_fraction = numeric.notna().mean() if n > 0 else 0

        # ----------------------------------------------------
        # Continuous/numeric ADMET
        # ----------------------------------------------------

        if n > 0 and numeric_fraction == 1:

            x = numeric

            results.append({
                "Variable": col,
                "Type": "Continuous",
                "N": n,
                "Missing": missing,
                "Unique": unique,
                "Mean": x.mean(),
                "SD": x.std(),
                "Median": x.median(),
                "Q1 (25%)": x.quantile(0.25),
                "Q3 (75%)": x.quantile(0.75),
                "IQR": x.quantile(0.75) - x.quantile(0.25),
                "Minimum": x.min(),
                "Maximum": x.max(),
                "Skewness": x.skew(),
                "Kurtosis": x.kurtosis()
            })

        # ----------------------------------------------------
        # Categorical ADMET
        # ----------------------------------------------------

        else:

            counts = non_missing.value_counts(dropna=False)

            # Overall summary row
            results.append({
                "Variable": col,
                "Type": "Categorical",
                "N": n,
                "Missing": missing,
                "Unique": unique,
                "Mean": np.nan,
                "SD": np.nan,
                "Median": np.nan,
                "Q1 (25%)": np.nan,
                "Q3 (75%)": np.nan,
                "IQR": np.nan,
                "Minimum": np.nan,
                "Maximum": np.nan,
                "Skewness": np.nan,
                "Kurtosis": np.nan
            })

    return pd.DataFrame(results)


admet_summary = admet_statistics(
    df,
    ADMET_COLS
)

# ============================================================
# 8. ADMET FREQUENCY DISTRIBUTIONS
# ============================================================

def categorical_frequency_table(data, columns):

    tables = []

    for col in columns:

        series = data[col]

        counts = series.value_counts(
            dropna=False
        )

        total = len(series)

        for category, count in counts.items():

            if pd.isna(category):
                category_label = "Missing"
            else:
                category_label = str(category)

            percentage = (count / total) * 100

            tables.append({
                "Variable": col,
                "Category": category_label,
                "Frequency": count,
                "Percentage": percentage
            })

    return pd.DataFrame(tables)


admet_frequencies = categorical_frequency_table(
    df,
    ADMET_COLS
)

# ============================================================
# 9. OVERALL LIBRARY SUMMARY
# ============================================================

overall_summary = pd.DataFrame({
    "Metric": [
        "Total compounds",
        "Total variables",
        "Physicochemical descriptors",
        "Bioactivity endpoints",
        "ADMET endpoints",
        "Morgan fingerprint size"
    ],
    "Value": [
        len(df),
        len(df.columns),
        len(PHYSICOCHEMICAL_COLS),
        len(BIOACTIVITY_COLS),
        len(ADMET_COLS),
        "1024-bit"
    ]
})

# ============================================================
# 10. MISSING DATA SUMMARY
# ============================================================

missing_summary = pd.DataFrame({
    "Variable": df.columns,
    "N": len(df),
    "Missing": df.isna().sum(),
    "Missing (%)": (
        df.isna().sum() / len(df) * 100
    ),
    "Complete (%)": (
        df.notna().sum() / len(df) * 100
    )
})

# Only keep variables relevant to the research question
missing_summary_research = missing_summary[
    missing_summary["Variable"].isin(
        PHYSICOCHEMICAL_COLS +
        BIOACTIVITY_COLS +
        ADMET_COLS
    )
].copy()

# ============================================================
# 11. DESCRIPTOR DISTRIBUTION SUMMARY
# ============================================================

# Useful compact table specifically for the thesis
distribution_summary = pd.concat([
    physchem_stats.assign(Category="Physicochemical"),
    bioactivity_stats.assign(Category="Predicted Bioactivity")
], ignore_index=True)

distribution_summary = distribution_summary[
    [
        "Category",
        "Variable",
        "N",
        "Missing",
        "Mean",
        "SD",
        "Median",
        "Q1 (25%)",
        "Q3 (75%)",
        "IQR",
        "Minimum",
        "Maximum",
        "Skewness",
        "Kurtosis"
    ]
]

# ============================================================
# 12. SAVE EVERYTHING TO EXCEL
# ============================================================

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    overall_summary.to_excel(
        writer,
        sheet_name="Library Summary",
        index=False
    )

    physchem_stats.to_excel(
        writer,
        sheet_name="Physicochemical",
        index=False
    )

    bioactivity_stats.to_excel(
        writer,
        sheet_name="Bioactivity",
        index=False
    )

    admet_summary.to_excel(
        writer,
        sheet_name="ADMET Summary",
        index=False
    )

    admet_frequencies.to_excel(
        writer,
        sheet_name="ADMET Frequencies",
        index=False
    )

    distribution_summary.to_excel(
        writer,
        sheet_name="Main Statistics",
        index=False
    )

    missing_summary_research.to_excel(
        writer,
        sheet_name="Missing Data",
        index=False
    )

# ============================================================
# 13. FORMAT EXCEL FILE
# ============================================================

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = load_workbook(OUTPUT_FILE)

for ws in wb.worksheets:

    # Header formatting
    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Freeze header
    ws.freeze_panes = "A2"

    # Auto-fit columns
    for column_cells in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[
            column_letter
        ].width = min(max_length + 2, 35)

    # Format numerical values
    for row in ws.iter_rows(
        min_row=2
    ):

        for cell in row:

            if isinstance(
                cell.value,
                float
            ):

                cell.number_format = "0.000"

wb.save(OUTPUT_FILE)

# ============================================================
# 14. PRINT RESULTS TO TERMINAL
# ============================================================

print("\n" + "=" * 70)
print("PHYSICOCHEMICAL DESCRIPTIVE STATISTICS")
print("=" * 70)

print(
    physchem_stats.to_string(
        index=False,
        float_format=lambda x: f"{x:.3f}"
    )
)

print("\n" + "=" * 70)
print("PREDICTED BIOACTIVITY DESCRIPTIVE STATISTICS")
print("=" * 70)

print(
    bioactivity_stats.to_string(
        index=False,
        float_format=lambda x: f"{x:.3f}"
    )
)

print("\n" + "=" * 70)
print("PREDICTED ADMET SUMMARY")
print("=" * 70)

print(
    admet_summary.to_string(
        index=False
    )
)

print("\n" + "=" * 70)
print("ADMET FREQUENCY DISTRIBUTIONS")
print("=" * 70)

print(
    admet_frequencies.to_string(
        index=False
    )
)

print("\n" + "=" * 70)
print(f"Statistics saved to: {OUTPUT_FILE}")
print("=" * 70)